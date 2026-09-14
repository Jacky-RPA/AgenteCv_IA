import re
from io import BytesIO

from docx import Document
from docx import shared as docx_shared
from docx.enum.text import WD_ALIGN_PARAGRAPH
from langchain_google_genai import ChatGoogleGenerativeAI

from dotenv import load_dotenv
import os

load_dotenv()

cv_llm = ChatGoogleGenerativeAI(
    model=os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite-preview"),
    temperature=0.2,
    google_api_key=os.getenv("GEMINI_API_KEY"),
    max_tokens=5000,
    timeout=60,
    max_retries=2,
)


def _contenido_de_respuesta(respuesta) -> str:
    contenido = getattr(respuesta, "content", "")
    if isinstance(contenido, str):
        return contenido.strip()
    if isinstance(contenido, list):
        partes = []
        for bloque in contenido:
            if isinstance(bloque, str):
                partes.append(bloque)
            elif isinstance(bloque, dict) and bloque.get("text"):
                partes.append(str(bloque["text"]))
        return "\n".join(partes).strip()
    return ""


def generar_cv_optimizado_para_vacante(
    cv_texto: str,
    oferta: str,
    informe: str = "",
) -> str:
    limpieza_prompt = f"""Eres un corrector profesional de CV. Limpia y corrige el CV original antes de adaptarlo a una vacante.

CV ORIGINAL:
{cv_texto}

Corrige ortografía, tildes, espacios, palabras pegadas, fechas mal separadas y errores de extracción PDF. Elimina solo duplicados exactos, símbolos extraños, Markdown, tablas rotas y frases que claramente no pertenecen a un CV. Conserva todos los datos reales: nombre, contacto, estudios, empresas, cargos, fechas, responsabilidades, logros y habilidades. No resumas, no inventes y devuelve únicamente el CV limpio en texto plano.
"""

    try:
        limpieza = _contenido_de_respuesta(cv_llm.invoke(limpieza_prompt))
    except Exception:
        limpieza = ""
    cv_limpio = limpieza or cv_texto.strip()

    prompt = f"""Eres un especialista senior en selección y adaptación de CV. Recibirás un CV ya corregido y una vacante. Debes adaptar el CV al puesto, no limitarte a copiarlo.

VACANTE / PUESTO AL QUE POSTULA:
{oferta}

CV LIMPIO DEL CANDIDATO:
{cv_limpio}

INFORME CONSOLIDADO DE COMPATIBILIDAD:
{informe or "Aún no hay informe disponible; analiza directamente la vacante y el CV."}

ETAPA 1 - LIMPIEZA Y CORRECCIÓN DEL CV ORIGINAL:
- Corrige ortografía, tildes, mayúsculas, espacios, palabras pegadas, fechas mal separadas y errores producidos al extraer texto de un PDF.
- Elimina duplicados exactos, caracteres sueltos, encabezados repetidos, tablas rotas, Markdown, emojis, símbolos extraños y frases que claramente no pertenecen al CV.
- Elimina frases genéricas o vacías que no aporten información profesional.
- Corrige la redacción para que suene natural, humana y profesional.
- No elimines información factual: empresas, cargos, fechas, estudios, responsabilidades, herramientas, metodologías ni logros.

ETAPA 2 - ADAPTACIÓN PROFUNDA A LA VACANTE:
- Analiza primero el cargo objetivo, responsabilidades, requisitos obligatorios, requisitos deseables, herramientas, competencias y palabras clave de la vacante.
- Construye una correspondencia entre cada prioridad del puesto y la evidencia concreta del CV limpio.
- Cambia el enfoque profesional del documento para que parezca preparado específicamente para esa postulación, no un CV genérico.
- Redacta PERFIL PROFESIONAL desde la perspectiva del puesto objetivo: menciona primero la experiencia real, los resultados y las herramientas que resuelven las necesidades de la vacante.
- Puedes cambiar la segunda línea por una denominación profesional objetivo coherente con la experiencia real, pero no presentes ese título como un cargo laboral que el candidato ya ocupó.
- Reordena las responsabilidades dentro de cada experiencia: primero las que más se relacionan con la vacante y después las demás. Conserva todas.
- Reordena HABILIDADES TÉCNICAS colocando primero las tecnologías y competencias reales que aparecen en la vacante, y luego el resto de habilidades del CV.
- Usa palabras clave de la vacante únicamente cuando estén respaldadas por el CV. Si una palabra clave no está respaldada, no la agregues como competencia.
- Usa el informe como diagnóstico: refuerza en el perfil y en el orden del CV lo que el informe clasifica como "Cumple" o "Cumple parcialmente" cuando exista evidencia real.
- Atiende las brechas del informe mediante mejor redacción, orden y claridad, pero no inventes información para ocultarlas.
- Si el puesto solicita algo que no aparece en el CV, no lo inventes: evita destacarlo y deja que la brecha sea visible en el informe.
- No conviertas una habilidad relacionada en experiencia profesional.
- No agregues competencias, años, cargos, empresas, logros o responsabilidades que no aparezcan en el CV original.

REGLAS OBLIGATORIAS DE CONSERVACIÓN:
- Conserva nombres, teléfonos, correos, enlaces, instituciones, empresas, cargos, fechas, responsabilidades, logros, herramientas y metodologías.
- No elimines ninguna experiencia, estudio, habilidad, logro ni categoría técnica.
- No cambies cargos reales, fechas ni nombres de empresa.
- No agregues competencias que no aparezcan explícitamente en el CV.
- Usa la vacante únicamente para enfocar el perfil profesional y ordenar las habilidades reales.
- En FORMACIÓN ACADÉMICA Y COMPLEMENTARIA y EXPERIENCIA PROFESIONAL conserva todo el contenido factual después de corregir su redacción y presentación. No resumas ni elimines contenido.

FORMATO OBLIGATORIO:
- Debes adaptar cualquier CV recibido exactamente al formato del CV de referencia de Jackelin Nuñez Aguirre.
- Devuelve únicamente el CV completo en texto plano. No uses Markdown, #, negritas, cursivas, tablas, emojis, iconos, separadores ni explicaciones.
- No agregues títulos como "CV Optimizado", "Resumen", "Competencias Clave" o "Puesto Objetivo".
- Usa exactamente estos cuatro títulos, en mayúsculas y en este orden: PERFIL PROFESIONAL, FORMACIÓN ACADÉMICA Y COMPLEMENTARIA, EXPERIENCIA PROFESIONAL, HABILIDADES TÉCNICAS.
- La primera línea debe ser el nombre completo real del candidato.
- La segunda línea debe ser el cargo real principal del candidato, no el nombre de la vacante.
- La tercera línea debe contener todos los datos de contacto disponibles, separados por " | ". No inventes ciudad, LinkedIn, portafolio ni datos faltantes.
- Conserva todas las experiencias en EXPERIENCIA PROFESIONAL y todos los estudios en FORMACIÓN ACADÉMICA Y COMPLEMENTARIA.
- En HABILIDADES TÉCNICAS conserva todas las categorías y tecnologías, pero puedes ordenar primero las que coincidan con la vacante.
- No incluyas una sección de referencias, objetivos, intereses, salario, disponibilidad o información personal si no aparece en el CV original.
- No incluyas instrucciones, notas del editor, advertencias ni explicaciones fuera del CV.

La salida debe tener esta forma exacta:

NOMBRE COMPLETO
CARGO REAL PRINCIPAL
TELÉFONO | CORREO | USUARIO | ENLACE
PERFIL PROFESIONAL
 Párrafo profesional de 6 a 8 líneas, completamente enfocado en el cargo objetivo y redactado con hechos, herramientas y logros del CV.
FORMACIÓN ACADÉMICA Y COMPLEMENTARIA
Todos los estudios y cursos del CV, una línea por elemento.
EXPERIENCIA PROFESIONAL
Todas las empresas, cargos, fechas, responsabilidades y logros del CV.
HABILIDADES TÉCNICAS
Todas las categorías y tecnologías del CV.

Antes de responder, comprueba que el perfil está claramente enfocado en el puesto objetivo, que las habilidades están priorizadas para la vacante, que las responsabilidades relevantes aparecen primero, que no inventaste datos y que no omitiste ningún bloque del CV limpio. Respeta exactamente los cuatro títulos y su orden.
"""

    try:
        respuesta = cv_llm.invoke(prompt)
        contenido = _contenido_de_respuesta(respuesta)
        if contenido:
            return contenido

        respuesta_reintento = cv_llm.invoke(
            f"{prompt}\nResponde ahora con el CV completo. No devuelvas una respuesta vacía."
        )
        contenido = _contenido_de_respuesta(respuesta_reintento)
        return contenido or cv_limpio
    except Exception:
        return cv_limpio


def _configurar_estilos_cv(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = docx_shared.Pt(12)
    normal.paragraph_format.space_after = 4
    normal.paragraph_format.line_spacing = 1.15

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.bold = True
        style.font.size = docx_shared.Pt(12)
        style.font.color.rgb = docx_shared.RGBColor(31, 78, 121)
        style.paragraph_format.keep_with_next = True

    styles["Heading 2"].paragraph_format.space_before = 12
    styles["Heading 2"].paragraph_format.space_after = 4
    styles["Heading 3"].font.color.rgb = docx_shared.RGBColor(55, 55, 55)
    styles["Heading 3"].paragraph_format.space_before = 6
    styles["Heading 3"].paragraph_format.space_after = 2


def _agregar_contenido_cv(doc: Document, contenido: str) -> None:
    secciones = {
        "PERFIL PROFESIONAL",
        "FORMACIÓN ACADÉMICA Y COMPLEMENTARIA",
        "FORMACION ACADEMICA Y COMPLEMENTARIA",
        "EXPERIENCIA PROFESIONAL",
        "HABILIDADES TÉCNICAS",
        "HABILIDADES TECNICAS",
    }
    encabezado = 0
    for linea in contenido.splitlines():
        texto = linea.strip()
        if not texto:
            continue
        texto = re.sub(r"^#{1,3}\s+", "", texto)
        texto = re.sub(r"\*\*(.*?)\*\*", r"\1", texto)
        texto = re.sub(r"\*(.*?)\*", r"\1", texto)
        texto = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", texto)
        texto = re.sub(r"[\U0001F300-\U0001FAFF\u2600-\u27BF]", "", texto).strip()
        if not texto or texto.upper().startswith("CV OPTIMIZADO PARA"):
            continue

        mayusculas = texto.upper()
        if encabezado == 0 and mayusculas not in secciones:
            parrafo = doc.add_heading(texto, level=1)
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in parrafo.runs:
                run.bold = True
            encabezado = 1
        elif mayusculas in secciones:
            parrafo = doc.add_heading(mayusculas, level=2)
            parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif re.match(r"^[•▪◦*-]\s*", texto):
            doc.add_paragraph(re.sub(r"^[•▪◦*-]\s*", "", texto), style="List Bullet")
        else:
            parrafo = doc.add_paragraph(texto)
            if encabezado == 1:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in parrafo.runs:
                    run.bold = True
                encabezado = 2
            elif encabezado == 2:
                parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in parrafo.runs:
                    run.font.color.rgb = docx_shared.RGBColor(90, 90, 90)
                encabezado = 3


def crear_word_cv_optimizado(
    cv_texto: str,
    oferta: str,
    informe: str = "",
) -> BytesIO:
    contenido = generar_cv_optimizado_para_vacante(cv_texto, oferta, informe)
    if not contenido.strip():
        contenido = cv_texto.strip()

    doc = Document()
    _configurar_estilos_cv(doc)
    section = doc.sections[0]
    section.top_margin = docx_shared.Inches(0.55)
    section.bottom_margin = docx_shared.Inches(0.55)
    section.left_margin = docx_shared.Inches(0.7)
    section.right_margin = docx_shared.Inches(0.7)
    _agregar_contenido_cv(doc, contenido)

    for parrafo in doc.paragraphs:
        for run in parrafo.runs:
            run.font.size = docx_shared.Pt(12)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
